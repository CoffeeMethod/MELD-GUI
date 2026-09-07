"""Pull plain text out of uploaded files.

MELD scores prose, so extraction aims for readable text with paragraph breaks
intact -- the model card asks for original formatting to be preserved. PDF and
DOCX support degrade gracefully: if the optional dependency is missing the user
gets a clear message instead of a stack trace.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json",
    ".log", ".text", ".org", ".tex",
}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}

SUPPORTED = TEXT_SUFFIXES | PDF_SUFFIXES | DOCX_SUFFIXES


class ExtractionError(RuntimeError):
    """Raised when a file cannot be turned into text."""


@dataclass(slots=True)
class Extracted:
    text: str
    filename: str
    kind: str
    pages: int = 0


def _suffix(filename: str) -> str:
    idx = filename.rfind(".")
    return filename[idx:].lower() if idx != -1 else ""


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _clean_pdf_text(text: str) -> str:
    """Undo the hard line wrapping most PDF extractors produce."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Rejoin words split across a line by a hyphen.
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    # A single newline inside a paragraph becomes a space; blank lines stay.
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
    return re.sub(r"[ \t]{2,}", " ", text).strip()


def extract(data: bytes, filename: str) -> Extracted:
    """Turn an uploaded file into text, or raise :class:`ExtractionError`."""
    suffix = _suffix(filename)

    if suffix in PDF_SUFFIXES:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover
            raise ExtractionError(
                "PDF support needs the 'pypdf' package (pip install pypdf)."
            ) from exc
        try:
            reader = PdfReader(io.BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:
            raise ExtractionError(f"Could not read the PDF: {exc}") from exc
        text = _clean_pdf_text("\n\n".join(pages))
        if not text.strip():
            raise ExtractionError(
                "No text found in this PDF. Scanned documents need OCR first."
            )
        return Extracted(text, filename, "pdf", pages=len(pages))

    if suffix in DOCX_SUFFIXES:
        try:
            import docx  # type: ignore[import-untyped]
        except ImportError as exc:  # pragma: no cover
            raise ExtractionError(
                "DOCX support needs the 'python-docx' package (pip install python-docx)."
            ) from exc
        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise ExtractionError(f"Could not read the DOCX file: {exc}") from exc
        blocks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.append(" ".join(cell.text for cell in row.cells))
        text = "\n\n".join(b.strip() for b in blocks if b.strip())
        if not text.strip():
            raise ExtractionError("The DOCX file contains no readable text.")
        return Extracted(text, filename, "docx")

    if suffix in TEXT_SUFFIXES or not suffix:
        text = _decode(data)
        if not text.strip():
            raise ExtractionError("The file is empty.")
        return Extracted(text, filename, "text")

    raise ExtractionError(
        f"Unsupported file type '{suffix}'. Supported: "
        + ", ".join(sorted(SUPPORTED))
    )
